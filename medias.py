# codigo para calcular a media de precipitacao por data a partir de uma planilha ja gerada

import sys

import openpyxl                  # para ler e gerar o excel
from openpyxl.styles import Font # para formatar o excel
from datetime import datetime    # para ordenar as linhas por data

ENTRADA = r"C:\Users\Milena\Desktop\dados\planilhas\saidaGEFS.xlsx"      # planilha já gerada (com várias linhas por data)
SAIDA = r"C:\Users\Milena\Desktop\dados\planilhas\saidaGEFS_medias.xlsx" # planilha nova, com as médias


# Função para ler a planilha de entrada e devolver as linhas de dados (sem cabeçalho):
def ler_planilha(caminho: str) -> list[list]:
    wb = openpyxl.load_workbook(caminho, data_only=True)
    ws = wb.active

    linhas = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0] is None:  # pula linhas vazias
            continue
        linhas.append(list(row))
    return linhas


# Função para agrupar as linhas por data e calcular a média da precipitação:
def calcular_medias_por_data(linhas: list[list]) -> list[list]:
    grupos: dict[str, list[float]] = {}
    info_por_data: dict[str, tuple] = {}

    for linha in linhas:
        data, usina, lat, lon, valor = linha
        try:
            valor_float = float(str(valor).replace(",", "."))
        except (TypeError, ValueError):
            print(f"[aviso] Valor inválido ignorado: '{valor}' (data {data})")
            continue
        grupos.setdefault(data, []).append(valor_float)
        info_por_data[data] = (usina, lat, lon)

    linhas_medias = []
    for data, valores in grupos.items():
        media = sum(valores) / len(valores)
        usina, lat, lon = info_por_data[data]
        linhas_medias.append([data, usina, lat, lon, round(media, 2)])

    # ordena por data (aceita tanto texto "dd/mm/aaaa" quanto datetime/date já lidos do excel)
    def chave_ordenacao(linha):
        data = linha[0]
        if isinstance(data, str):
            return datetime.strptime(data, "%d/%m/%Y")
        return data

    linhas_medias.sort(key=chave_ordenacao)
    return linhas_medias


# Função para gerar a planilha de saída com as médias:
def gerar_excel(linhas: list[list], saida: str) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Medias"

    cabecalho = ["Data", "Usina", "Latitude", "Longitude", "Preciptação [mm]"]
    ws.append(cabecalho)

    for celula in ws[1]:
        celula.font = Font(name="Arial", bold=True)

    for linha in linhas:
        ws.append(linha)

    for row in ws.iter_rows(min_row=2):
        for celula in row:
            celula.font = Font(name="Arial")

    larguras = [12, 12, 14, 14, 18]
    for i, largura in enumerate(larguras, start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = largura

    wb.save(saida)
    print(f"Excel gerado em: {saida}")


def main():
    if len(sys.argv) >= 3:
        entrada = sys.argv[1]
        saida = sys.argv[2]
    else:
        entrada = ENTRADA
        saida = SAIDA

    linhas = ler_planilha(entrada)
    if not linhas:
        print("Nenhuma linha encontrada na planilha de entrada.")
        return

    linhas_medias = calcular_medias_por_data(linhas)
    gerar_excel(linhas_medias, saida)


if __name__ == "__main__":
    main()