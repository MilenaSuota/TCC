# codigo para extrair dados do arquivo psath compactado

import re       # para leitura do titulo do arquivo
import sys      # para pegar argumentos da linha de comando
import zipfile  # para extrair arquivos zip

from datetime import datetime, timedelta  # para ordenar as linhas por data e somar dias
from pathlib import Path                  # para manipular os caminhos dos arquivos

import openpyxl                  # para gerar o excel
from openpyxl.styles import Font # para formatar o excel


ZIP = r"C:\Users\laura\Documents\UTFPR\TCC\GEFS50_precipitacao14d_20260922.zip"              # caminho do zip
ARQUIVO_EXTRAIDO = r"C:\Users\laura\Documents\UTFPR\TCC\dados_extraidos"    # pasta onde os dat serão extraídos
SAIDA = r"C:\Users\laura\Documents\UTFPR\TCC\saidaGEFS.xlsx"                    # excel gerado
USINA = "PSATMAU"                                                           # usina

# Funcao para extrair o zip e retornar o path da pasta extraida:
def extrair_zip(caminho_zip: str, pasta_saida: str) -> Path:
    pasta_saida = Path(pasta_saida)
    pasta_saida.mkdir(parents=True, exist_ok=True)

    with zipfile.ZipFile(caminho_zip, "r") as arquivo_zip:
        arquivo_zip.extractall(pasta_saida)

    return pasta_saida

# Funcao para extrair a data do nome do arquivo (formato antigo, DDMMAAAA):
def extrair_data_do_nome(nome_arquivo: str) -> str | None:
    match = re.search(r"(\d{2})(\d{2})(\d{4})", nome_arquivo)
    if not match:
        return None
    dia, mes, ano = match.groups()
    return f"{dia}/{mes}/{ano}"

# Funcao para extrair a data base do nome do arquivo GEFS_m_* (formato DDMMAA):
def extrair_data_base_serie(nome_arquivo: str) -> datetime | None:
    match = re.search(r"(\d{2})(\d{2})(\d{2})(?!\d)", nome_arquivo)
    if not match:
        return None
    dia, mes, ano = match.groups()
    return datetime(2000 + int(ano), int(mes), int(dia))

# Função para encontrar a linha que contém a usina de Maua:
def encontrar_linha_palavra(linhas: list[str], palavra: str) -> str | None:
    padrao = re.compile(rf"\b{re.escape(palavra)}\b")
    for linha in linhas:
        if padrao.search(linha):
            return linha
    return None

# Função para separar os campos da linha encontrada:
def separar_campos(linha: str) -> list[str]:
    campos = linha.split()
    if len(campos) >= 4:
        return campos

    for sep in [";", ","]:
        campos = [c.strip() for c in linha.split(sep)]
        if len(campos) >= 4:
            return campos
    return campos  # devolve o que tiver, mesmo que incompleto

# Função para processar arquivos do tipo GEFS_m_* (uma linha por usina, várias colunas de dias):
def processar_arquivo_serie(caminho: Path, palavra: str) -> list[list[str]]:
    data_base = extrair_data_base_serie(caminho.name)
    if not data_base:
        print(f"[aviso] Não consegui extrair a data base de: {caminho.name}")
        return []

    try:
        with open(caminho, "r", encoding="latin-1") as f:
            linhas = f.readlines()
    except Exception as e:
        print(f"[erro] Falha ao ler {caminho.name}: {e}")
        return []

    linha_alvo = encontrar_linha_palavra(linhas, palavra)
    if not linha_alvo:
        print(f"[aviso] Linha com '{palavra}' não encontrada em {caminho.name}")
        return []

    campos = separar_campos(linha_alvo)
    if len(campos) < 4:
        print(f"[aviso] Linha em {caminho.name} não tem campos suficientes: {campos}")
        return []

    latitude, longitude = campos[1], campos[2]
    valores = campos[3:]  # os valores de precipitação, um por dia

    saida = []
    for i, valor in enumerate(valores, start=1):
        data = data_base + timedelta(days=i)
        saida.append([data.strftime("%d/%m/%Y"), palavra, latitude, longitude, valor])
    return saida

# Função principal que processa a pasta e monta as linhas de saída:
def processar_pasta(pasta: Path, palavra: str) -> list[list[str]]:
    linhas_saida = []
    arquivos_dat = sorted(pasta.rglob("*.dat"))
    if not arquivos_dat:
        print(f"[aviso] Nenhum arquivo .dat encontrado em {pasta}")

    for arquivo in arquivos_dat:
        # arquivos de série (várias datas na mesma linha)
        if "GEFS_m_" in arquivo.name:
            linhas_saida.extend(processar_arquivo_serie(arquivo, palavra))
            continue

        # arquivos antigos (uma data por arquivo, nome com 8 dígitos)
        data_formatada = extrair_data_do_nome(arquivo.name)
        if not data_formatada:
            print(f"[aviso] Não consegui extrair a data do nome: {arquivo.name}")
            continue
        try:
            with open(arquivo, "r", encoding="latin-1") as f:
                linhas = f.readlines()
        except Exception as e:
            print(f"[erro] Falha ao ler {arquivo.name}: {e}")
            continue
        linha_alvo = encontrar_linha_palavra(linhas, palavra)
        if not linha_alvo:
            print(f"[aviso] Linha com '{palavra}' não encontrada em {arquivo.name}")
            continue
        campos = separar_campos(linha_alvo)
        if len(campos) < 4:
            print(f"[aviso] Linha encontrada em {arquivo.name} não tem campos "
                  f"suficientes: {campos}")
            continue
        latitude = campos[1]
        longitude = campos[2]
        precipitacao = campos[3]

        linhas_saida.append([data_formatada, palavra, latitude, longitude, precipitacao])

    linhas_saida.sort(key=lambda linha: datetime.strptime(linha[0], "%d/%m/%Y"))
    return linhas_saida

# Função para gerar o Excel a partir das linhas coletadas:
def gerar_excel(linhas: list[list[str]], saida: str) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "PSATMAU"

    cabecalho = ["Data", "Usina", "Latitude", "Longitude", "Preciptação [mm]"]
    ws.append(cabecalho)

    for celula in ws[1]:
        celula.font = Font(name="Arial", bold=True)

    for linha in linhas:
        ws.append(linha)

    for row in ws.iter_rows(min_row=2):
        for celula in row:
            celula.font = Font(name="Arial")
    larguras = [12, 12, 14, 14, 18]    # ajusta largura das colunas
    for i, largura in enumerate(larguras, start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = largura

    wb.save(saida)
    print(f"Excel gerado em: {saida}")

# Função principal:
def main():
    if len(sys.argv) >= 3:
        zip = sys.argv[1]
        saida = sys.argv[2]
    else:
        zip = ZIP
        saida = SAIDA

    pasta = extrair_zip(zip, ARQUIVO_EXTRAIDO)
    linhas = processar_pasta(pasta, USINA)

    if not linhas:
        print("Nenhuma linha válida encontrada. Excel não será gerado.")
        return
    gerar_excel(linhas, saida)

if __name__ == "__main__":
    main()
